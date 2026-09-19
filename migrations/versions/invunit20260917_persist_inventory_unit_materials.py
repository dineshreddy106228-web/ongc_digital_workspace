"""Persist per-material DFS/ST assignments for multi-unit inventory assets.

Revision ID: invunit20260917
Revises: qcctrl20260901_sap_docs
"""

from io import BytesIO
import re

from alembic import op
from openpyxl import load_workbook
import sqlalchemy as sa


revision = "invunit20260917"
down_revision = "qcctrl20260901_sap_docs"
branch_labels = None
depends_on = None


def _text(value):
    if value is None:
        return ""
    return str(value).strip()


def _material_code(value):
    text = _text(value)
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    if text.isdigit() and len(text) < 9:
        text = text.zfill(9)
    return text


def upgrade():
    op.create_table(
        "inventory_monitoring_work_center_unit_materials",
        sa.Column("id", sa.BigInteger(), autoincrement=True, nullable=False),
        sa.Column("work_center_id", sa.BigInteger(), nullable=False),
        sa.Column("unit_type", sa.String(length=80), nullable=False),
        sa.Column("material_code", sa.String(length=64), nullable=False),
        sa.Column("mapping_batch_id", sa.BigInteger(), nullable=False),
        sa.Column("is_current", sa.Boolean(), nullable=False, server_default=sa.true()),
        sa.ForeignKeyConstraint(
            ["work_center_id"], ["inventory_monitoring_work_centers.id"], ondelete="CASCADE",
        ),
        sa.ForeignKeyConstraint(
            ["mapping_batch_id"], ["inventory_monitoring_upload_batches.id"], ondelete="CASCADE",
        ),
        sa.PrimaryKeyConstraint("id"),
        sa.UniqueConstraint(
            "work_center_id", "unit_type", "material_code", "mapping_batch_id",
            name="uq_inventory_monitoring_unit_material_version",
        ),
    )
    op.create_index(
        "ix_inventory_monitoring_unit_material_current",
        "inventory_monitoring_work_center_unit_materials",
        ["work_center_id", "is_current", "unit_type"],
    )

    # The original list is retained in the latest mapping batch on existing
    # installations.  Materialise it now so later rollback-payload retention
    # can safely purge the workbook bytes without losing DFS/ST classification.
    bind = op.get_bind()
    batch = bind.execute(sa.text(
        "SELECT id, source_data FROM inventory_monitoring_upload_batches "
        "WHERE source_group = 'mapping' AND source_file_size > 0 "
        "ORDER BY id DESC LIMIT 1"
    )).mappings().first()
    if not batch or not batch["source_data"]:
        return

    centres = {
        row["normalized_name"]: row["id"]
        for row in bind.execute(sa.text(
            "SELECT id, normalized_name FROM inventory_monitoring_work_centers"
        )).mappings()
    }
    workbook = load_workbook(BytesIO(bytes(batch["source_data"])), read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    next(rows, None)
    seen = set()
    payload = []
    for row in rows:
        if len(row) < 4:
            continue
        centre_name, unit_type = _text(row[2]), _text(row[3])
        centre_id = centres.get(re.sub(r"\s+", " ", centre_name).casefold())
        if not centre_id or not unit_type:
            continue
        for value in row[4:]:
            code = _material_code(value)
            key = (centre_id, unit_type, code)
            if not code or key in seen:
                continue
            seen.add(key)
            payload.append({
                "work_center_id": centre_id,
                "unit_type": unit_type,
                "material_code": code,
                "mapping_batch_id": batch["id"],
            })
    if payload:
        table = sa.table(
            "inventory_monitoring_work_center_unit_materials",
            sa.column("work_center_id", sa.BigInteger()),
            sa.column("unit_type", sa.String()),
            sa.column("material_code", sa.String()),
            sa.column("mapping_batch_id", sa.BigInteger()),
        )
        op.bulk_insert(table, payload)


def downgrade():
    op.drop_index(
        "ix_inventory_monitoring_unit_material_current",
        table_name="inventory_monitoring_work_center_unit_materials",
    )
    op.drop_table("inventory_monitoring_work_center_unit_materials")
